"""Protocol JSON (contract ``protocol``), the Приложение 2 golden strings, DOCX / HTML / PDF views."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import pytest

from inspector_common.contracts.loader import validation_errors
from inspector_common.params import load_params
from inspector_compare.engine import compare_object
from inspector_compare.export.submission import build_submission
from inspector_compare.protocol import texts as T
from inspector_compare.protocol.builder import (
    ProtocolBuilder,
    build_protocol,
    content_sha256,
    passport_from_context,
)

# ── Приложение 2 golden strings ───────────────────────────────────────────────────────────────

SAMPLE_SECTION2 = [
    ("Всего параметров в Матрице", 132, "100%"),
    ("Проверено успешно (есть ПД, РД, ИД)", 112, "84,8%"),
    ("Не проверено (отсутствует ИД)", 18, "13,6%"),
    ("Не загружены документы (технические ошибки)", 2, "1,6%"),
    ("Выявлено нарушений (всего)", 12, "9,1%"),
    ("─ Критических (приостановка)", 6, "4,5%"),
    ("─ Существенных (предписание)", 6, "4,5%"),
    ("Подозрений ИИ (свободный поиск)", 3, "2,3%"),
]


def test_percent_rules_reproduce_the_sample() -> None:
    assert [T.percent_text(p) for p in T.partition_percents([112, 18, 2], 132)] == ["84,8%", "13,6%", "1,6%"]
    assert T.percent_text(T.percent_half_up(12, 132)) == "9,1%"
    assert T.percent_text(T.percent_half_up(6, 132)) == "4,5%"
    assert T.percent_text(T.percent_half_up(3, 132)) == "2,3%"
    assert T.percent_text(T.percent_half_up(132, 132)) == "100%"
    assert T.percent_text(T.percent_half_up(0, 132)) == "0%"
    parts = T.partition_percents([2, 14, 2, 114], 132)
    assert sum(parts) == 100
    assert T.ru_date(dt.date(2026, 7, 1)) == "1 июля 2026 г."


def test_section2_reproduces_the_sample_table(tmp_path: Path, syn: Any) -> None:
    ctx = syn.make_context(tmp_path, [])
    codes = [p.code for p in sorted(load_params(), key=lambda s: s.param_id)]
    buckets = ["CHECKED_OK"] * 112 + ["NOT_CHECKED_NO_ID"] * 18 + ["NOT_LOADED_TECH_ERRORS"] * 2
    rows = [
        {
            "parameter_code": code,
            "location": "OBJECT",
            "violation_label": "NO_VIOLATION",
            "protocol_status": "OK",
            "decision_trace": {"bucket": b},
        }
        for code, b in zip(codes, buckets, strict=True)
    ]
    builder = ProtocolBuilder(ctx, [], rows, {"checks": []}, run_id="t")
    builder.scenario = "FULL"
    s2 = builder.section2(6, 6, 3)
    assert [(r["label_ru"], r["count"], r["percent_text"]) for r in s2["rows"]] == SAMPLE_SECTION2
    assert s2["footnotes"][-1] == T.DISCLAIMER


def test_scenario_line_is_russian_only() -> None:
    assert T.scenario_line("PD_RD_ONLY", "Только ПД и РД") == "Тип проверки: Только ПД и РД"
    assert (
        T.scenario_line("FULL", "Полный комплект (ПД, РД, ИД)")
        == "Тип проверки: Полный комплект (ПД, РД, ИД)"
    )


def test_equal_counts_get_equal_percents_and_the_adjustment_is_named() -> None:
    parts, adjusted = T.partition_percents_ex([2, 13, 2, 115], 132)
    assert parts[0] == parts[2]
    assert sum(parts) == 100
    assert adjusted == [1]  # the group of two equal rows never takes a lone 0,1


def test_stage_and_basis_labels() -> None:
    assert T.basis_ru("MANDATORY_SOURCE_MISSING") == "Отсутствует обязательный источник данных"
    assert T.file_stage_ru(None, "UNKNOWN") == T.SERVICE_STAGE_RU
    assert T.file_stage_ru(None, "RD_ID_MIXED") == "РД/ИД (смешанный)"
    assert T.file_stage_ru("RD", "RD_ID_MIXED") == "РД"
    status, stage, basis = T.a1_cells(
        {"protocol_status": "COMPARISON_IMPOSSIBLE", "completeness_basis": "VALUE_ABSTAINED"}
    )
    assert (status, stage) == ("Проверка невозможна", T.NO_MISSING_STAGE_RU)
    assert basis == "Значение параметра не извлечено"


def test_passport_comes_from_the_title_blocks(tmp_path: Path, syn: Any) -> None:
    """«Объект:» / «Адрес:» from the object name printed most often in the title blocks (AG-02B), split at the
    address; the manifest corpus name when no title block carries one."""
    printed = "Школа на 600 мест, р-н Богородское, ул. Тюменская, влд. 5"

    def with_names(*names: str) -> Any:
        titles = [dict(syn.title(10 + i, i + 1, "План"), object_name=n) for i, n in enumerate(names)]
        return syn.make_context(tmp_path / str(len(names)), [syn.layout("F9002", "RD", titles=titles)])

    passport = passport_from_context(with_names(printed, printed, "Школа на 600 мест"))
    assert (passport.name, passport.address) == (
        "Школа на 600 мест",
        "р-н Богородское, ул. Тюменская, влд. 5",
    )
    assert passport_from_context(with_names()).name == "Синтетический объект"
    p = passport_from_context(with_names("Жилой дом по адресу: г. Москва, ул. Речная, д. 7"))
    assert (p.name, p.address) == ("Жилой дом", "г. Москва, ул. Речная, д. 7")
    assert passport_from_context(with_names("Жилой дом с подземной автостоянкой")).address is None


# ── full build ────────────────────────────────────────────────────────────────────────────────


def _value(syn: Any, code: str, stage: str, file_id: str, value: Any, unit: str) -> dict[str, Any]:
    return {
        "value_id": f"v-{code}-{stage}",
        "object_id": syn.OBJ,
        "file_id": file_id,
        "file_sha256": syn.sha(int(file_id[1:])),
        "stage": stage,
        "page_no": 3,
        "param_code": code,
        "value_raw": str(value),
        "value_norm": {"type": "number", "value": value, "unit": unit},
        "method": "TABLE",
        "confidence": 0.95,
        "quality_flag": "OK",
        "pipeline_version": "test",
    }


@pytest.fixture()
def protocol_inputs(tmp_path: Path, syn: Any, cfg: Any) -> dict[str, Any]:
    values = [
        _value(syn, "SPZU-025", "PD", "F9001", 850, "м²"),
        _value(syn, "SPZU-025", "RD", "F9002", 720, "м²"),
    ]
    ctx = syn.make_context(tmp_path, syn.ventilation_layouts(), values=values)
    result = compare_object(ctx, cfg, run_id="t-run")
    sub = build_submission(ctx.object_id, result.findings)
    return {"ctx": ctx, "groups": result.groups, "findings": result.findings, "submission": sub.extended}


def test_protocol_validates_and_follows_appendix2(protocol_inputs: dict[str, Any]) -> None:
    p = build_protocol(
        protocol_inputs["ctx"],
        protocol_inputs["groups"],
        protocol_inputs["findings"],
        protocol_inputs["submission"],
        run_id="t-run",
        generated_at=dt.datetime(2026, 7, 1, 10, 0, tzinfo=dt.UTC),
    )
    assert validation_errors("protocol", p) == []
    h = p["header"]
    assert h["title"] == "ПРОТОКОЛ АВТОМАТИЗИРОВАННОЙ СВЕРКИ № 2026-07-01-SYN-1-1"
    assert h["generated_at_ru"] == "1 июля 2026 г."
    assert h["version_line"] == "1 (предварительная)"
    assert h["status_line"] == "⚠️ ОЖИДАЕТ ВЕРИФИКАЦИИ (дозагрузка возможна)"
    a2 = p["appendix2"]
    s4, s5 = a2["section4_critical"], a2["section5_substantial"]
    assert [r["no"] for r in s4["rows"]] == list(range(1, s4["count"] + 1))
    assert [r["no"] for r in s5["rows"]] == list(
        range(s4["count"] + 1, s4["count"] + s5["count"] + 1)
    )  # continuous
    assert s5["rows"][0]["parameter_code"] == "SPZU-025"
    assert s5["rows"][0]["pd"] == "850 м²" and s5["rows"][0]["rd"] == "720 м²" and s5["rows"][0]["id"] == "—"
    assert s5["rows"][0]["deviation"]["text"].startswith("⬇️ Уменьшение на 15,3%")
    miss = s4["rows"][0]
    assert miss["parameter_label"].startswith(
        "Воздуховоды общеобменной вентиляции (IOS4-078), пом. 101 [карточка Б."
    )
    assert miss["inspector_decision_ru"] == "⏳ Ожидает"
    s7 = a2["section7_resolution"]
    assert len(s7["critical"]) == s4["count"] and len(s7["substantial"]) == s5["count"]
    assert all(
        r["recommendation"].startswith("Проект рекомендации (до подтверждения инспектором): ")
        for r in s7["critical"] + s7["substantial"]
    )
    assert s7["critical"][0]["work_type"].endswith("(IOS4-078)")
    assert s7["substantial"][0]["violation_kind"].endswith("(SPZU-025)")
    s2 = {r["key"]: r for r in a2["section2_summary"]["rows"]}
    assert s2["TOTAL_PARAMS"]["count"] == 132
    partition = [
        r
        for r in a2["section2_summary"]["rows"]
        if r["key"] in ("CHECKED_OK", "NOT_CHECKED_NO_ID", "NOT_LOADED_TECH_ERRORS", "NOT_CHECKED_NO_PD_RD")
    ]
    assert sum(r["count"] for r in partition) == 132
    assert round(sum(r["percent"] for r in partition), 1) == 100.0
    assert s2["VIOLATIONS_TOTAL"]["count"] == s4["count"] + s5["count"]
    cards = p["evidence_cards"]
    assert [c["card_no"] for c in cards] == [f"Б.{n}" for n in range(1, len(cards) + 1)]
    assert all(c["sources"] for c in cards)
    reg = p["input_registry"]["files"]
    assert {r["file_id"] for r in reg if r["used"]} == {"F9001", "F9002"}
    assert next(r for r in reg if r["file_id"] == "F9006")["exclusion_reason"]
    assert p["submission_checks"] == protocol_inputs["submission"]["checks"]


def test_content_hash_ignores_the_date(protocol_inputs: dict[str, Any]) -> None:
    args = (
        protocol_inputs["ctx"],
        protocol_inputs["groups"],
        protocol_inputs["findings"],
        protocol_inputs["submission"],
    )
    a = build_protocol(*args, run_id="r1", generated_at=dt.datetime(2026, 7, 1, tzinfo=dt.UTC))
    b = build_protocol(*args, run_id="r2", generated_at=dt.datetime(2026, 8, 2, tzinfo=dt.UTC))
    assert a["content_sha256"] == b["content_sha256"] == content_sha256(a)
    assert a["protocol_no"] != b["protocol_no"]


def test_docx_html_and_pdf_views(tmp_path: Path, protocol_inputs: dict[str, Any]) -> None:
    from docx import Document

    from inspector_compare.protocol.docx_render import render_docx
    from inspector_compare.protocol.html_render import render_html
    from inspector_compare.protocol.pdf import render_pdf

    p = build_protocol(
        protocol_inputs["ctx"],
        protocol_inputs["groups"],
        protocol_inputs["findings"],
        protocol_inputs["submission"],
        run_id="t",
    )
    path = render_docx(p, tmp_path / "p.docx")
    doc = Document(str(path))
    section = doc.sections[0]
    margins = (
        section.left_margin.mm,
        section.right_margin.mm,
        section.top_margin.mm,
        section.bottom_margin.mm,
    )
    assert [round(m, 1) for m in margins] == [30.0, 15.0, 20.0, 20.0]  # ГОСТ margins of the Приложение 2 docx
    assert round(section.page_width.mm) == 210 and round(section.page_height.mm) == 297
    assert doc.paragraphs[0].text == p["header"]["title"]
    assert doc.paragraphs[0].runs[0].font.name == "Segoe UI" and doc.paragraphs[0].runs[0].bold
    text = "\n".join(par.text for par in doc.paragraphs)
    for n in (1, 2, 7):
        assert T.SECTION_TITLES[n] in text
    assert "7.1. По критическим нарушениям" in text and "ПРИЛОЖЕНИЕ Б. КАРТОЧКИ ДОКАЗАТЕЛЬСТВ" in text
    assert len(doc.tables) >= 8
    html = render_html(p)
    assert "None" not in html and "{" not in html.split("<style>")[0] + html.split("</style>")[1]
    for n in (1, 2, 7):
        assert T.SECTION_TITLES[n] in html
    pdf = render_pdf(html, tmp_path / "p.pdf", engine="pymupdf")
    assert pdf.pages >= 3 and pdf.engine == "pymupdf"
    import pymupdf

    with pymupdf.open(pdf.path) as d:
        assert "РАЗДЕЛ 1" in d[0].get_text()
    # Evidence cards carry every ТЗ §9.2 source field: full SHA-256, шифр, редакция, approval status, regions.
    sha = p["evidence_cards"][0]["sources"][0]["file_sha256"]
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for view in (html, cells):
        assert f"SHA-256 {sha}" in view
        assert "статус утверждения: нет данных" in view and "области (" in view and "шифр " in view
    assert "ПД (ожидаемое): F9001" in cells and "РД (фактическое): F9002" in cells
    # An empty section says so (the substantial sections have rows here; 6 has none without AG-07's hook).
    assert p["appendix2"]["section6_ai_suspicions"]["count"] == 0
    assert T.EMPTY_SECTION in html and T.EMPTY_SECTION in text


@pytest.mark.slow
def test_chromium_pdf_when_available(tmp_path: Path, protocol_inputs: dict[str, Any]) -> None:
    from inspector_compare.protocol.html_render import render_html
    from inspector_compare.protocol.pdf import find_chromium, render_pdf

    if find_chromium() is None:
        pytest.skip("Chromium/Chrome is not installed")
    p = build_protocol(
        protocol_inputs["ctx"],
        protocol_inputs["groups"],
        protocol_inputs["findings"],
        protocol_inputs["submission"],
        run_id="t",
    )
    pdf = render_pdf(render_html(p), tmp_path / "p.pdf", engine="chromium")
    assert pdf.engine == "chromium" and pdf.pages >= 3
    import pymupdf

    with pymupdf.open(pdf.path) as d:
        assert abs(d[0].rect.width - 595) < 2 and abs(d[0].rect.height - 842) < 2  # A4 portrait
        assert "РАЗДЕЛ 1. СТАТУС ЗАГРУЗКИ ДОКУМЕНТОВ" in d[0].get_text()


def test_card_thumbnails_are_framed_crops_embedded_in_the_views(
    tmp_path: Path, protocol_inputs: dict[str, Any]
) -> None:
    """93 §5.5: evidence cards carry page crops with a blue frame (expected) and a red one (actual); DOCX gets
    pictures, the HTML (Chromium PDF) data URIs, the PyMuPDF fallback drops them; a missing file is skipped."""
    import io

    import pymupdf
    from docx import Document
    from PIL import Image

    from inspector_compare.protocol.docx_render import render_docx
    from inspector_compare.protocol.html_render import render_html
    from inspector_compare.protocol.pdf import render_pdf
    from inspector_compare.protocol.thumbs import BLUE, RED, card_thumbnails

    pdfs: dict[str, Path] = {}
    for file_id, pages in (("F9001", 10), ("F9002", 7)):
        doc = pymupdf.open()
        for n in range(pages):
            page = doc.new_page(width=842, height=595)
            page.insert_text((60, 80), f"{file_id} page {n + 1}", fontsize=18)
        pdfs[file_id] = tmp_path / f"{file_id}.pdf"
        doc.save(pdfs[file_id])
        doc.close()
    p = build_protocol(
        protocol_inputs["ctx"],
        protocol_inputs["groups"],
        protocol_inputs["findings"],
        protocol_inputs["submission"],
        run_id="t",
    )
    thumbs = card_thumbnails(p, pdfs.get)
    assert thumbs, "every card cites F9001/F9002 pages"
    cards = {c["card_no"]: c for c in p["evidence_cards"]}
    colours = set()
    for items in thumbs.values():
        for t in items:
            image = Image.open(io.BytesIO(t.png)).convert("RGB")
            assert image.width == t.width == 900
            found = {c for _, c in image.getcolors(maxcolors=1 << 24) or []}
            colour = BLUE if t.role in ("EXPECTED", "SUPPORTING_EXPECTED") else RED
            assert ("синяя рамка" in t.caption) == (colour == BLUE)
            boxes = (cards[t.card_no]["sources"][t.source_index].get("geometry") or {}).get("boxes")
            # framed when the source has regions; a value without geometry shows the whole page, unframed
            assert (colour in found) == bool(boxes), (t.card_no, t.role)
            if boxes:
                colours.add(colour)
    assert colours == {BLUE, RED}
    total = sum(len(v) for v in thumbs.values())
    docx_path = render_docx(p, tmp_path / "t.docx", thumbs)
    assert len(Document(str(docx_path)).inline_shapes) == total
    html = render_html(p, thumbs)
    assert html.count("data:image/png;base64,") == total
    fallback = render_pdf(html, tmp_path / "t.pdf", engine="pymupdf")
    assert fallback.pages >= 3
    # a page that cannot be located is skipped, never an error
    assert card_thumbnails(p, lambda file_id: None) == {}


def test_value_conflict_suspicions_reach_section6_cards_and_a5_without_touching_violations(
    tmp_path: Path, syn: Any, cfg: Any
) -> None:
    def fire(stage: str, file_id: str, value: str) -> dict[str, Any]:
        v = _value(syn, "PZ-022", stage, file_id, value, "")
        v["value_norm"] = {"type": "enum", "value": value}
        v["value_id"] = f"v-PZ-022-{file_id}"
        return v

    values = [fire("PD", "F9001", "II"), fire("PD", "F9004", "III")]
    ctx = syn.make_context(tmp_path, syn.ventilation_layouts(), values=values)
    ctx.run_dir = tmp_path / "run"
    result = compare_object(ctx, cfg, run_id="t-run")
    sub = build_submission(ctx.object_id, result.findings)
    kwargs = {"run_id": "t-run", "generated_at": dt.datetime(2026, 7, 1, 10, 0, tzinfo=dt.UTC)}
    p = build_protocol(ctx, result.groups, result.findings, sub.extended, **kwargs)
    ctx.run_dir = None  # no hypothesis run: the same protocol without the suspicion
    base = build_protocol(ctx, result.groups, result.findings, sub.extended, **kwargs)
    assert validation_errors("protocol", p) == []
    s6, s6_base = p["appendix2"]["section6_ai_suspicions"], base["appendix2"]["section6_ai_suspicions"]
    assert s6["count"] == 1 and s6_base["count"] == 0
    rows = s6["rows"]
    assert rows[0]["method"] == "LOGICAL_ANALYSIS" and rows[0].get("parameter_code") is None
    assert rows[0]["description"].startswith("В ПД противоречивые значения параметра PZ-022")
    card = next(c for c in p["evidence_cards"] if c["card_no"] == rows[0]["card_ref"])
    assert card["parameter_code"] == "PZ-022" and {s["file_id"] for s in card["sources"]} == {
        "F9001",
        "F9004",
    }
    assert any(r["card_ref"] == rows[0]["card_ref"] for r in p["tz92_tables"]["a5_hypotheses"])
    # the card shows the conflicting values per stage and per page, and the catalog criticality
    assert card["stage_values"]["RD"] is None and "файл" in card["stage_values"]["PD"]
    assert {s["file_id"]: s.get("value_text") for s in card["sources"]} == {"F9001": "II", "F9004": "III"}
    assert card["criticality"] is not None and all("role" not in s for s in card["sources"])
    # violations, labels and the submission are exactly those of the protocol without the suspicion
    assert p["submission_checks"] == base["submission_checks"]
    assert p["appendix2"]["section4_critical"] == base["appendix2"]["section4_critical"]
    assert p["appendix2"]["section5_substantial"] == base["appendix2"]["section5_substantial"]
    s2 = {r["key"]: r["count"] for r in p["appendix2"]["section2_summary"]["rows"]}
    s2_base = {r["key"]: r["count"] for r in base["appendix2"]["section2_summary"]["rows"]}
    assert s2["AI_SUSPICIONS"] == 1 and s2["VIOLATIONS_TOTAL"] == s2_base["VIOLATIONS_TOTAL"]
