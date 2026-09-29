"""Protocol JSON → DOCX in the Приложение 2 look (93 §5.2, measured from the sample docx).

A4 portrait (11906×16838 twips); margins top 20 mm, bottom 20 mm, left 30 mm, right 15 mm; font «Segoe UI»
(declared; Selawik is the metric-compatible substitute on hosts without it), colour #0F1115; title and
«РАЗДЕЛ N» headings 16.5 pt bold with 24 pt before / 12 pt after; «7.1./7.2.» 15 pt bold; header block 12 pt with
bold labels and soft line breaks; tables 11.5 pt, single 0.5 pt borders on all edges, header row bold and repeated.
Emoji status markers are kept as text (Word draws them with the system emoji font).
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from inspector_compare.protocol import texts as T

FONT = "Segoe UI"
INK = RGBColor(0x0F, 0x11, 0x15)
MUTED = RGBColor(0x55, 0x5B, 0x66)


def _set_font(run: Any, size: float, bold: bool = False, italic: bool = False, color: RGBColor = INK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


# Cell padding (top, bottom, left, right) in dxa: the Приложение 2 docx pads body cells 150/150/0/240; the left 0
# is kept at 100 so text does not touch the border. Appendix tables (А.1 lists every parameter) are compact.
BODY_PADDING = (150, 150, 100, 240)
COMPACT_PADDING = (60, 60, 100, 100)


class DocxWriter:
    def __init__(self) -> None:
        self.padding: tuple[int, int, int, int] = BODY_PADDING
        self.doc = Document()
        section = self.doc.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(12)
        normal.font.color.rgb = INK
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), FONT)
        pf = normal.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)

    # ── paragraphs ────────────────────────────────────────────────────────────────────────────

    def heading(self, text: str, size: float = 16.5, before: float = 24, after: float = 12) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_with_next = True
        _set_font(p.add_run(text), size, bold=True)

    def labelled_lines(self, pairs: Sequence[tuple[str, str]]) -> None:
        """One paragraph, «Label: value» lines separated by soft line breaks (the sample's header block)."""
        p = self.doc.add_paragraph()
        for i, (label, value) in enumerate(pairs):
            if i:
                p.add_run().add_break(WD_BREAK.LINE)
            _set_font(p.add_run(label), 12, bold=True)
            _set_font(p.add_run(" " + value), 12)

    def text(
        self, value: str, size: float = 12, italic: bool = False, color: RGBColor = INK, before: float = 6
    ) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        _set_font(p.add_run(value), size, italic=italic, color=color)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── tables ────────────────────────────────────────────────────────────────────────────────

    def table(self, header: Sequence[str], rows: Iterable[Sequence[Any]], size: float = 11.5) -> None:
        """A bordered table; rows are appended one by one (``table.cell(i, j)`` is quadratic in python-docx)."""
        tbl = self.doc.add_table(rows=1, cols=len(header))
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        self._borders(tbl)
        for cell, h in zip(tbl.rows[0].cells, header, strict=True):
            self._cell(cell, str(h), size, bold=True)
        tr_pr = tbl.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)
        for row in rows:
            for cell, value in zip(tbl.add_row().cells, row, strict=True):
                self._cell(cell, "" if value is None else str(value), size)
        self.doc.add_paragraph()

    def key_value_table(self, pairs: Sequence[tuple[str, Any]], size: float = 10.5) -> None:
        tbl = self.doc.add_table(rows=0, cols=2)
        self._borders(tbl)
        for k, v in pairs:
            key_cell, value_cell = tbl.add_row().cells
            self._cell(key_cell, k, size, bold=True)
            self._cell(value_cell, "—" if v in (None, "", []) else str(v), size)
        self.doc.add_paragraph()

    def _cell(self, cell: Any, value: str, size: float, bold: bool = False) -> None:
        p = cell.paragraphs[0]
        p.clear()
        p.paragraph_format.space_after = Pt(0)
        _set_font(p.add_run(value), size, bold=bold)
        tc_pr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for side, width in zip(("top", "bottom", "left", "right"), self.padding, strict=True):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(width))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tc_pr.append(mar)

    def thumbnails(self, thumbs: Sequence[Any], width_mm: float = 78, max_height_mm: float = 95) -> None:
        """Two thumbnails per row (no borders), each with its caption (blue frame = expected, red = actual)."""
        cols = 2
        tbl = self.doc.add_table(rows=0, cols=cols)
        for start in range(0, len(thumbs), cols):
            cells = tbl.add_row().cells
            for cell, t in zip(cells, thumbs[start : start + cols], strict=False):
                p = cell.paragraphs[0]
                p.clear()
                p.paragraph_format.space_after = Pt(2)
                height = width_mm * t.height / max(1, t.width)
                width = width_mm if height <= max_height_mm else max_height_mm * t.width / max(1, t.height)
                p.add_run().add_picture(io.BytesIO(t.png), width=Mm(width))
                _set_font(cell.add_paragraph().add_run(t.caption), 8.5, color=MUTED)
        self.doc.add_paragraph()

    @staticmethod
    def _borders(tbl: Any) -> None:
        tbl_pr = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "auto")
            borders.append(el)
        tbl_pr.append(borders)

    def save(self, path: Path, title: str) -> Path:
        props = self.doc.core_properties
        props.title = title
        props.subject = "Протокол автоматизированной сверки (Приложение 2)"
        props.author = "Инспектор ИИ"
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(".docx.tmp")
        self.doc.save(str(tmp))
        tmp.replace(path)
        return path


def _or(value: Any, default: str = "нет данных") -> str:
    return default if value in (None, "") else str(value)


def render_docx(
    protocol: dict[str, Any], path: Path, thumbnails: Mapping[str, Sequence[Any]] | None = None
) -> Path:
    """The protocol as DOCX in the Приложение 2 look; ``thumbnails`` (card_no → Thumbnail) go under their cards."""
    w = DocxWriter()
    h = protocol["header"]
    obj = protocol["object"]
    L = T.HEADER_LABELS
    w.heading(h["title"], before=24, after=12)
    w.labelled_lines([(L["object"], _or(obj.get("name"))), (L["address"], _or(obj.get("address")))])
    pairs = [
        (L["case"], _or(obj.get("supervision_case_no"))),
        (L["customer"], _or(obj.get("customer"))),
        (L["contractor"], _or(obj.get("contractor"))),
        (L["date"], h["generated_at_ru"]),
        (L["version"], h["version_line"]),
        (L["status"], h["status_line"] + (T.RECHECK_SUFFIX if h.get("recheck_running") else "")),
    ]
    if h.get("scenario_line"):
        pairs.append((L["scenario"], h["scenario_line"].removeprefix(L["scenario"]).strip()))
    w.labelled_lines(pairs)

    a2 = protocol["appendix2"]
    # Раздел 1
    w.heading(T.SECTION_TITLES[1])
    s1 = a2["section1_load_status"]
    w.table(
        T.COLUMNS[1],
        [
            (
                r["stage_ru"],
                r["status_ru"],
                r["files_loaded_text"],
                _or(r.get("files_expected"), "—"),
                r["comment"],
            )
            for r in s1["rows"]
        ],
    )
    if s1.get("scenario_line"):
        w.text(s1["scenario_line"], size=11.5, before=0)
    # Раздел 2
    w.heading(T.SECTION_TITLES[2])
    s2 = a2["section2_summary"]
    w.table(T.COLUMNS[2], [(r["label_ru"], r["count"], r["percent_text"]) for r in s2["rows"]])
    for note in s2.get("footnotes") or []:
        w.text("* " + note, size=10, italic=True, color=MUTED, before=2)
    # Раздел 3
    s3 = a2["section3_not_checked_no_id"]
    w.heading(T.SECTION_TITLES[3].format(n=s3["count"]))
    if not s3["rows"]:
        w.text(T.EMPTY_SECTION, before=0)
    else:
        w.table(
            T.COLUMNS[3],
            [
                (r["no"], r["parameter_code"], r["section_ru"], r["parameter_name"], r["missing_document"])
                for r in s3["rows"]
            ],
        )
    # Разделы 4–5
    for key, num in (("section4_critical", 4), ("section5_substantial", 5)):
        sec = a2[key]
        w.heading(T.SECTION_TITLES[num].format(n=sec["count"]))
        if not sec["rows"]:
            w.text(T.EMPTY_SECTION, before=0)
        else:
            w.table(
                T.COLUMNS[4],
                [
                    (
                        r["no"],
                        r["section_ru"],
                        r["parameter_label"],
                        r["pd"],
                        r["rd"],
                        r["id"],
                        r["deviation"]["text"],
                        r["inspector_decision_ru"],
                    )
                    for r in sec["rows"]
                ],
                size=10.5,
            )
    # Раздел 6
    s6 = a2["section6_ai_suspicions"]
    w.heading(T.SECTION_TITLES[6].format(n=s6["count"]))
    if not s6["rows"]:
        w.text(T.EMPTY_SECTION, before=0)
    else:
        w.table(
            T.COLUMNS[6],
            [
                (
                    r["no"],
                    r["method_ru"],
                    r["description"] + f" [карточка {r['card_ref']}]",
                    r["pd"],
                    r["rd"],
                    r["id"],
                    r["inspector_decision_ru"],
                    r["rejection_reason"],
                    r["ai_comment"],
                )
                for r in s6["rows"]
            ],
            size=9.5,
        )
    # Раздел 7
    s7 = a2["section7_resolution"]
    w.heading(T.SECTION_TITLES[7])
    w.heading(T.SUBSECTION_71, size=15)
    if not s7["critical"]:
        w.text(T.EMPTY_SECTION, before=0)
    else:
        w.table(T.COLUMNS[71], [(r["no"], r["work_type"], r["recommendation"]) for r in s7["critical"]])
    w.heading(T.SUBSECTION_72, size=15)
    if not s7["substantial"]:
        w.text(T.EMPTY_SECTION, before=0)
    else:
        w.table(
            T.COLUMNS[72], [(r["no"], r["violation_kind"], r["recommendation"]) for r in s7["substantial"]]
        )

    # Приложение А
    tz = protocol["tz92_tables"]
    w.padding = COMPACT_PADDING
    w.page_break()
    w.heading(T.APPENDIX_TITLES["A"], before=0)
    w.text(tz.get("banner_ru") or T.BANNER_TZ92, size=11, italic=True, color=MUTED, before=0)
    w.heading(T.APPENDIX_TITLES["A1"], size=13, before=12, after=6)
    if tz.get("a1_completeness"):
        w.table(
            (
                "Код",
                "Параметр",
                "Статус",
                "Отсутствует стадия",
                "Документ",
                "Комплектность / основание",
                "Действие",
            ),
            [
                (
                    r["parameter_code"],
                    r["parameter_name"],
                    T.a1_cells(r)[0],
                    T.a1_cells(r)[1],
                    _or(r.get("document"), "—"),
                    T.a1_cells(r)[2],
                    _or(r.get("action_ru"), "—"),
                )
                for r in tz["a1_completeness"]
            ],
            size=8.5,
        )
    w.heading(T.APPENDIX_TITLES["A2"], size=13, before=12, after=6)
    if tz.get("a2_candidates"):
        w.table(
            (
                "Карточка",
                "Параметр",
                "Места",
                "Ожидаемое",
                "Фактическое",
                "Отклонение",
                "Риск",
                "Уверенность",
                "Решение",
            ),
            [
                (
                    r["card_ref"],
                    r["parameter_label"],
                    ", ".join(r.get("locations") or []),
                    _or(r.get("expected"), "—"),
                    _or(r.get("actual"), "—"),
                    _or(r.get("deviation_text"), "—"),
                    T.enum_ru("RiskLevel", r.get("risk_level")),
                    f"{float(r.get('confidence') or 0):.2f}".replace(".", ","),
                    T.DECISION_RU.get(r["inspector_status"], r["inspector_status"]),
                )
                for r in tz["a2_candidates"]
            ],
            size=8.5,
        )
    w.heading(T.APPENDIX_TITLES["A3"], size=13, before=12, after=6)
    w.text(
        "Решений инспектора нет (протокол предварительный)." if not tz.get("a3_confirmed") else "",
        size=11,
        before=0,
    )
    w.heading(T.APPENDIX_TITLES["A4"], size=13, before=12, after=6)
    if tz.get("a4_negative_verified"):
        w.table(
            ("Параметр", "Кем", "Основание"),
            [
                (
                    r["parameter_label"],
                    "система" if r["decided_by"] == "SYSTEM" else "инспектор",
                    _or(r.get("reason_ru"), "—"),
                )
                for r in tz["a4_negative_verified"]
            ],
            size=9.5,
        )
    else:
        w.text("Нет.", size=11, before=0)
    w.heading(T.APPENDIX_TITLES["A5"], size=13, before=12, after=6)
    if tz.get("a5_hypotheses"):
        w.table(
            ("Карточка", "Код", "Описание", "Уверенность", "Нормативная база", "Привязка", "Решение"),
            [
                (
                    r["card_ref"],
                    _or(r.get("parameter_code"), "—"),
                    r["description"],
                    f"{float(r.get('confidence') or 0):.2f}".replace(".", ","),
                    _or(r.get("normative_base"), "—"),
                    _or(r.get("evidence_bind_status"), "—"),
                    T.DECISION_RU.get(r["inspector_status"], r["inspector_status"]),
                )
                for r in tz["a5_hypotheses"]
            ],
            size=8.5,
        )

    # Приложение Б
    w.page_break()
    w.heading(T.APPENDIX_TITLES["B"], before=0)
    for card in protocol.get("evidence_cards") or []:
        w.heading(f"Карточка {card['card_no']}. {card['parameter_label']}", size=12.5, before=12, after=4)
        sources = T.source_lines(card.get("sources"))
        w.key_value_table(
            [
                ("Идентификаторы находок", ", ".join(card.get("finding_ids") or [])),
                ("Код параметра", card["parameter_code"]),
                ("Места", ", ".join(card.get("locations") or [])),
                ("Ожидаемое значение", card.get("expected_value")),
                ("Фактическое значение", card.get("actual_value")),
                ("Критичность", card.get("criticality")),
                (
                    "Уровень риска / приоритет",
                    T.risk_priority(card.get("risk_level"), card.get("review_priority")),
                ),
                ("Источники", sources),
                ("Обоснование", card.get("rationale")),
                ("Версия правила", card.get("rule_version")),
                ("Согласованное изменение", card.get("approved_change_ref") or "не найдено"),
                (
                    "Решение инспектора",
                    T.DECISION_RU.get(card["inspector"]["status"], card["inspector"]["status"]),
                ),
            ]
        )
        if thumbnails and thumbnails.get(card["card_no"]):
            w.thumbnails(thumbnails[card["card_no"]])

    # Приложение В
    w.page_break()
    w.heading(T.APPENDIX_TITLES["V"], before=0)
    reg = protocol["input_registry"]
    w.table(
        ("Файл", "Наименование", "Стадия", "Раздел", "Стр.", "SHA-256", "Использован", "Исключение"),
        [
            (
                r["file_id"],
                _or(r.get("file_name"), "—"),
                T.file_stage_ru(r.get("stage"), r.get("manifest_stage")),
                T.section_ru(r.get("section")),
                _or(r.get("pages"), "—"),
                str(r["sha256"])[:12] + "…",
                "да" if r["used"] else "нет",
                _or(r.get("exclusion_reason"), "—"),
            )
            for r in reg["files"]
        ],
        size=7.5,
    )
    v = protocol["versions"]
    w.key_value_table(
        [
            ("Версия конвейера", v.get("pipeline_version")),
            ("Версия Матрицы", v.get("matrix_version")),
            ("Версия правил сравнения", v.get("model_version")),
            ("Версия кода", v.get("code_version")),
            ("Версия контрактов", v.get("contract_version")),
            ("Хеш входного манифеста (input_manifest_hash)", protocol["input_manifest_hash"]),
            ("Хеш содержания протокола (content_sha256)", protocol.get("content_sha256")),
        ]
    )
    w.text("Инспектор: ______________________ / ______________________ /", size=12, before=24)
    w.text("Дата: «____» ______________ 20__ г.", size=12, before=6)
    return w.save(path, h["title"])
